#!/usr/bin/env python3
"""Example usage of Document & Excel Processing Toolkit."""

import pandas as pd
from pathlib import Path

def example_analyze_excel():
    """Example: Analyze an existing Excel file."""
    # Replace with your actual file path
    file_path = "your_file.xlsx"
    
    if Path(file_path).exists():
        # Read the Excel file
        df = pd.read_excel(file_path)
        
        # Basic analysis
        print(f"Shape: {df.shape}")
        print(f"Columns: {list(df.columns)}")
        print(f"\nFirst 5 rows:")
        print(df.head())
        
        # Statistical summary for numeric columns
        numeric_cols = df.select_dtypes(include=['number'])
        if not numeric_cols.empty:
            print(f"\nNumeric Summary:")
            print(numeric_cols.describe())
    else:
        print(f"File not found: {file_path}")

def example_create_excel():
    """Example: Create a new Excel file."""
    # Sample data
    data = {
        'Product': ['Apple', 'Banana', 'Cherry', 'Date'],
        'Price': [1.5, 0.8, 2.0, 3.5],
        'Quantity': [100, 150, 80, 60]
    }
    
    df = pd.DataFrame(data)
    output_file = "fruits_sales.xlsx"
    
    # Save to Excel
    df.to_excel(output_file, index=False)
    print(f"Created {output_file} with {len(df)} rows")
    
    # Verify
    df_read = pd.read_excel(output_file)
    print(f"Verified: {len(df_read)} rows read back")

def example_extract_pdf_text():
    """Example: Extract text from PDF (requires pdfplumber)."""
    try:
        import pdfplumber
        
        pdf_path = "your_document.pdf"
        if Path(pdf_path).exists():
            with pdfplumber.open(pdf_path) as pdf:
                text = ""
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
                
                print(f"Extracted {len(text)} characters from {pdf_path}")
                print("First 500 characters:")
                print(text[:500])
        else:
            print(f"PDF not found: {pdf_path}")
    except ImportError:
        print("pdfplumber not installed. Run: pip install pdfplumber")

def example_process_word_document():
    """Example: Process Word document (requires python-docx)."""
    try:
        from docx import Document
        
        docx_path = "your_document.docx"
        if Path(docx_path).exists():
            doc = Document(docx_path)
            
            # Extract paragraphs
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            print(f"Found {len(paragraphs)} paragraphs")
            
            # Extract tables
            tables = doc.tables
            print(f"Found {len(tables)} tables")
            
            if tables:
                print("First table preview:")
                for i, table in enumerate(tables[:1]):
                    for row in table.rows[:3]:  # First 3 rows
                        cells = [cell.text for cell in row.cells]
                        print(f"  {' | '.join(cells)}")
        else:
            print(f"Word document not found: {docx_path}")
    except ImportError:
        print("python-docx not installed. Run: pip install python-docx")

if __name__ == "__main__":
    print("Document & Excel Processing Toolkit - Examples")
    print("=" * 50)
    
    print("\n1. Excel Analysis Example:")
    example_analyze_excel()
    
    print("\n2. Excel Creation Example:")
    example_create_excel()
    
    print("\n3. PDF Text Extraction Example:")
    example_extract_pdf_text()
    
    print("\n4. Word Document Processing Example:")
    example_process_word_document()